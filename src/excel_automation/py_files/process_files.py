import os
import openpyxl
import re
import config as cfg
import streamlit as st
import fnmatch

from unidecode import unidecode
from stop_words import get_stop_words
from typing import Literal
from pathlib import Path
from .stlt_helpers import col_to_num
from .db_related import insert_catalogue_rows, keyword_search, get_db_columns
from .semantic_search import semantic_search


# regarding search file

@st.cache_resource
def column_names() -> tuple[list[str], int]:
   cols = get_db_columns("catalogue_items")
   shortened_cols = [col for col in cols if col != "id"]
   return shortened_cols, len(shortened_cols)


def insert_matches(matches: list, wb: openpyxl.Workbook, item: dict) -> None:
	ws = wb[item["sheet"]]
	row = item["row"]
	col = item["result_col"]

	if not matches:
		ws.cell(row=row, column=col).value = "Found nothing" # type: ignore (the cells are very likely non-merged)
		return

	cn, length = column_names()

	for match in matches:
		# 1. Resolve the lookup mechanism ONCE per row/match item outside the loop
		if isinstance(match, dict):
			def get_val(i: int): return match[cn[i]]
		elif isinstance(match, list):
			def get_val(i: int): return match[i]
		else:
			continue

		# 2. The inner loop stays pure, fast, and free of conditional overhead
		for i in range(length):
			incr_col: int = col + i
			ws.cell(row=row, column=incr_col).value = f"{get_val(i)}"  # type: ignore

		col += length + 1

	return


def search_for_keywords(payload: dict) -> tuple[bool, str]:

	for path, list in payload.items():
			wb = list[0]["wb"]
			for item in list:
				matches: list = keyword_search(item["kws"])
				matches.append(semantic_search(item["texts"]))
				insert_matches(matches, wb, item)
			wb.save(path)
			wb.close()
	
	success = f"Successfully inserted matches into {len(payload)} file/s."
	return True, success


def extract_keywords(cellvalue: str) -> list[str]:
	"""
	Extracts text from a cell and splits it into keywords, including abbreviations, anglicizing for optimal db search etc.
	"""
	stop_words = get_stop_words('spanish')
	words = [unidecode(w.lower()) for w in cellvalue.split()]
	separated_kws = [
		w.strip('.,;:!?()"\'') 
		for w in words 
		if w.strip('.,;:!?()"\'') and w.lower() not in stop_words
	]

	for idx, kw in enumerate(separated_kws):
		if kw in cfg.ABBREVIATIONS:
			kw_extended = cfg.ABBREVIATIONS[kw]
			separated_kws[idx] = kw_extended.lower()

	return separated_kws


# general/shared process

def clean_description(description: str) -> str:
	# standardizing descriptions
	preprocessed_description: str = description.strip().lower()
	cleaned_description: list = re.findall(r'\w+', preprocessed_description)
	for idx, item in enumerate(cleaned_description):
		if item in cfg.ABBREVIATIONS:
			cleaned_description[idx] = cfg.ABBREVIATIONS[item]

	return " ".join(cleaned_description)


def find_min_max_cols(list_of_cols: list) -> tuple:
	all_cols: set = set()
	for c_int in list_of_cols:
		all_cols.add(c_int)
	return min(all_cols), max(all_cols)


def processing_cols(item: dict, cols_list: list[int]) -> tuple[int, int, int]:
	if isinstance(item["min_col"], str) and item["min_col"].isalpha():
		item["min_col"] = min_col = col_to_num(item["min_col"], 1)
		item["max_col"] = max_col = col_to_num(item["max_col"], 1) if item["max_col"] is not None else item["min_col"]
	else:
		min_col = int(item["min_col"])
		max_col = int(item["max_col"]) if item["max_col"] is not None else min_col
	cols_list.extend([min_col, max_col])
	min_col, max_col = find_min_max_cols(cols_list)
	col_diff = max_col - min_col + 1

	return min_col, max_col, col_diff


def preparing_filters_and_cols(path_dict: dict) -> tuple[list, int, int, int]: 
	# pre-processing filters
	cols_for_minmax = []
	compiled_filters = []

	# parsing columns to 1-based indexing
	for f in path_dict["filters"]:

		raw_cols: str = str(f.get("targ_col_raw", ""))
		col_add: str | None = None
		if raw_cols.startswith("any"):
			col_add = "any"
		elif raw_cols.startswith("all"):
			col_add = "all"

		extracted_raw_cols: list = re.findall(r'[A-Z]+', raw_cols)
		cols_parsed = [col_to_num(col, 1) for col in extracted_raw_cols] # first is to find the min and the max of the columns, therefore 1-based indexing
		f.update({"col_add": col_add, "cols_parsed": cols_parsed})

		cols_for_minmax.extend(cols_parsed)

	min_col, max_col, col_diff = processing_cols(path_dict, cols_for_minmax)

	# preparing columns to 0-based indexing
	for f in path_dict["filters"]:

		omit_row = str(f.get("omit_row_when", ""))
		if omit_row:               
			delimiter = ";" if ";" in omit_row else ","
			targets: list = [
				v.strip().lower() 
				for v in omit_row.split(delimiter)
				if v.strip()
				]
		else:
			targets = [""]

		cols_0idx = [col - min_col for col in f["cols_parsed"]]
			
		compiled_filters.append({
			"col_add": f["col_add"],
			"columns": cols_0idx,
			"targets": targets
		})

	return compiled_filters, min_col, max_col, col_diff


def process_spreadsheet_files(action: Literal["catalogue", "search"], parameters: dict, db_path: Path | str | None = cfg.DB_PATH) -> tuple[bool, str]:

	if not parameters or db_path is None:
		print("Error: parameters and/or db_path is None")
		return False, ""

	#print(f"Loading data from {parameters} ... \n")

	# Initialize payload
	sPayload = {} # search payload
	cPayload = [] # catalogue payload
	
	# Load catalogues workbooks into memory
	for path, item in parameters.items():
		if not os.path.exists(path):
			print(f"Error: File '{path}' not found!")
			return False, ""
			
		wb_cat = openpyxl.load_workbook(path, data_only=True)

		# processing filters and columns
		compiled_filters, min_col, max_col, col_diff = preparing_filters_and_cols(item)

		# setting memory for results/matches column
		if action == "search":
			global rCol, targeted_cols
			rCol = col_to_num(item["results_col"], 1)
			sPayload[path] = []
			targeted_cols = (item["min_col"] - min_col, item["max_col"] - min_col)

		# processing sheets
		for sheet in wb_cat.worksheets:

			# Ignores hidden or leftover sheets
			if sheet.sheet_state != 'visible': continue
			
			#print(f"\n Processing sheet: {sheet.title}")
			if sheet.title not in item["selected_sheets"]: continue

			# Do NOT use values_only=True so we can access cell.coordinate
			for row in sheet.iter_rows(min_row=item["min_row"], max_row=item["max_row"], min_col=min_col, max_col=max_col):

				vals: list = [
					str(c.value).strip().lower() 
					if c.value is not None else "" 
					for c in row
					]

				# failsafe for missing columns
				if len(vals) < col_diff:
					for _ in range(len(vals), col_diff):
						vals.append("")

				# use of filter conditions
				skip_row = False
				for f in compiled_filters:
					col_add: str | None = f["col_add"]
					cols: list = f["columns"]
					targets: list = f["targets"]

					if cols:

						if col_add == "all":
							# Fix 3: Standardized fnmatch across all branches
							if all(any(fnmatch.fnmatch(vals[c], t) for t in targets) for c in cols):
								skip_row = True
								break

						elif col_add == "any":
							if any(any(fnmatch.fnmatch(vals[c], t) for t in targets) for c in cols):
								skip_row = True
								break

						else:
							if any(fnmatch.fnmatch(vals[cols[0]], t) for t in targets):
								skip_row = True
								break

					else:
						# --- CASE B: Check ANY Column in Row ---                        
						if any(
							fnmatch.fnmatch(cell_str, t) 
							for cell_str in vals 
							for t in targets
						):
							skip_row = True
							break

				if skip_row:
					continue  # SKIPS TO THE NEXT ROW IN SHEET.ITER_ROWS()

				if action == "search":

					collected_kws = []
					cell_text = ""
					for cell in vals[targeted_cols[0] : targeted_cols[1] + 1]:
						if cell:
							cell_text += f"{cell} "
							cur_kws: list = extract_keywords(cell)
							collected_kws.extend(cur_kws)
					sPayload[path].append({"kws": collected_kws, "texts": cell_text, "wb": wb_cat, "sheet": sheet.title, "result_col": rCol, "row": row[0].row})

				else:
					
					cm: list = item.get("column_mappings", [])
					row_data: dict = {}
					for mItem in cm:
						delimiter = mItem.get("delimiter")
						
						for k, v in mItem.items():
							if k == "delimiter":
								continue
								
							if delimiter:
								iter_vals = [vals[col_idx] for col_idx in v]
								row_data[k] = delimiter.join(iter_vals)
							else:
								col_idx = v[0]
								row_data[k] = vals[col_idx]

					# adding extra informations to the row_data
					#row_data.update({"sheet": sheet.title, "cell_ref": f"=[{path}]{sheet.title}!{row[0].coordinate}"})
					row_data.update({"sheet": sheet.title, "cell_ref": f"'[{os.path.basename(path)}]{sheet.title}!{row[0].coordinate}"})

					# adding row_data to the payload
					cPayload.append(row_data)

		wb_cat.close()

	if action == "catalogue":
		# Save data to database
		result, success = insert_catalogue_rows(cPayload, st.session_state["db_template"])
	else:
		# technical search and semantic search
		result, success = search_for_keywords(sPayload)
		if not result:
			print("Something went wrong with keyword search.")

	return result, success

