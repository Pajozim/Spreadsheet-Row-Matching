
import os
import openpyxl
import re
import unicodedata
import config as cfg
from pyxlsb import open_workbook as open_xlsb
from pathlib import Path
from excel_automation.py_files.db_related import keyword_search
from excel_automation.py_files.process_catalogues import clean_unit


def insert_matches(matches: list[dict], wb: openpyxl.Workbook, sheet: str, col: int, row: int) -> None:
    ws = wb[sheet]

    if not matches:
        ws.cell(row=row, column=col).value = "Found nothing" # type: ignore (the cells are very likely non-merged)
        return

    for match in matches:
        ws.cell(row=row, column=col).value = match["clave"]
        ws.cell(row=row, column=col + 1).value = match["description"]
        ws.cell(row=row, column=col + 2).value = match["unit"]
        ws.cell(row=row, column=col + 3).value = match["cell_reference"]
        col += 5

    return


def extract_keywords(cellvalue: str) -> list[str]:
    """
    Extracts text from a cell and splits it into keywords.
    """

    keywords = []

    if re.search(r'[.,]', cellvalue): 
        keywords = re.split(r'[.,]', cellvalue)[0]
    else:
        keywords = cellvalue

    return keywords.lower().split()


def process_sheets(wb: openpyxl.Workbook, wb_op: openpyxl.Workbook):
    for sheet in wb.worksheets:
        print(f"Processing sheet: {sheet.title}")

        # Ignores hidden or leftover sheets
        if sheet.sheet_state != 'visible': continue 

        # to focus the search under headers
        context = []
        activity = []
        details = []

        for row in sheet.iter_rows(min_row=3, min_col=1, max_col=13):

            if any(not cell for cell in row): continue

            #if any(getattr(cell, "is_merged", True) for cell in row):

            #else:

            # setting indexes
            partida_idx = 0
            helper1_idx = 10
            description_idx = 11
            unit_idx = 12

            # setting values
            pre_partida = row[partida_idx].value
            pre_helper1 = row[helper1_idx].value
            pre_description = row[description_idx].value
            pre_unit = row[unit_idx].value

            # Safe extraction + stripping whitespace
            partida = str(pre_partida).strip() if pre_partida is not None else ""
            helper1 = str(pre_helper1).strip() if pre_helper1 is not None else "" # helper for conditions
            #helper2 = str(row[2].value).strip() if row[2].value is not None else "" # helper for conditions
            description = str(pre_description).strip() if pre_description is not None else "" # description
            unit = str(pre_unit).strip() if pre_unit is not None else "" # unit column

            # Set context
            if helper1 and description and not unit:
                # specific context (partida/concepto)
                cleaned_context: list = re.findall(r'\w+', description)
                deaccented_context: list = [unicodedata.normalize("NFD", str(cc)).encode("ascii", "ignore").decode("utf-8") for cc in cleaned_context]
                only_3ormore: list = [cc for cc in deaccented_context if len(cc) >= 3]
                #trimmed_concept = " OR ".join([word.lower() + "*" for word in cleaned_concept if len(word) >= 3])

                # context packed
                for c in only_3ormore:
                    key = c.lower()
                    val = cfg.CONTEXT_ABB.get(key, f"{key}*")
                    context.append(val)

            if not description: continue

            kws: list[str] = extract_keywords(description)
            #print("kws:", kws)

            matches: list[dict] = keyword_search(kws, clean_unit(unit), partida, context)
            print(f"matches: {matches}")

            row_number: int = row[0].row # type: ignore
            if matches: insert_matches(matches, wb_op, sheet.title, 15, row_number)


def process_source_file(file_path: Path | str = cfg.SOURCE_FILE, output_path: Path | str = cfg.OUTPUT_FILE) -> None:

    if not file_path or not os.path.exists(file_path):
        print("Error: No file path provided / file not found.")
        return

    print(f"Loading source file {file_path} from .xlsx_files/ ... \n")


    wb = openpyxl.load_workbook(file_path, data_only=True, rich_text=True)
    wb_op = openpyxl.load_workbook(output_path, data_only=True)

    process_sheets(wb, wb_op)

    wb_op.save(output_path)
    wb_op.close()
    wb.close()

    return